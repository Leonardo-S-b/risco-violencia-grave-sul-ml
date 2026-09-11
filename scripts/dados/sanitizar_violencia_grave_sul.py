from __future__ import annotations

import csv
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PASTA_DADOS_BRUTOS = Path("data/raw")
PASTA_PROCESSADOS = Path("data/processed")
PASTA_RELATORIOS = Path("reports/dados")

ARQUIVO_BASE_LONGA = PASTA_PROCESSADOS / "violencia_grave_sul_municipal_longa.csv"
ARQUIVO_BASE_MENSAL = PASTA_PROCESSADOS / "violencia_grave_sul_municipal_mensal.csv"
RELATORIO = PASTA_RELATORIOS / "resumo_sanitizacao_violencia_grave_sul.txt"

ANO_INICIAL = 2015
ANO_FINAL = 2025

MAPA_UF_SUL = {
    "RS": "RS",
    "SC": "SC",
    "PR": "PR",
    "RIO GRANDE DO SUL": "RS",
    "SANTA CATARINA": "SC",
    "PARANA": "PR",
}

EVENTOS_VIOLENCIA_GRAVE = {
    "HOMICIDIO DOLOSO": "homicidio_doloso",
    "ROUBO SEGUIDO DE MORTE (LATROCINIO)": "latrocinio",
    "TENTATIVA DE HOMICIDIO": "tentativa_homicidio",
    "LESAO CORPORAL SEGUIDA DE MORTE": "lesao_corporal_seguida_morte",
    "FEMINICIDIO": "feminicidio",
}

CABECALHO_BASE_LONGA = [
    "ano_arquivo",
    "uf",
    "municipio",
    "evento_original",
    "evento",
    "data_referencia",
    "ano",
    "mes",
    "total_vitima",
    "abrangencia",
]

CABECALHO_BASE_MENSAL = [
    "uf",
    "municipio",
    "data_referencia",
    "ano",
    "mes",
    "total_violencia_grave",
    "homicidio_doloso",
    "latrocinio",
    "tentativa_homicidio",
    "lesao_corporal_seguida_morte",
    "feminicidio",
]


def normalizar_texto(valor: object) -> str:
    if valor is None:
        return ""
    texto = str(valor).strip().upper()
    texto = unicodedata.normalize("NFKD", texto)
    texto = texto.encode("ascii", "ignore").decode("ascii")
    return " ".join(texto.split())


def limpar_texto(valor: object) -> str:
    if valor is None:
        return ""
    return " ".join(str(valor).strip().split())


def pegar_ano_pelo_nome(caminho: Path) -> int:
    texto_ano = caminho.stem.split("_")[-1]
    if not texto_ano.isdigit():
        raise ValueError(f"Nao foi possivel identificar o ano: {caminho.name}")
    return int(texto_ano)


def converter_data(valor: object) -> tuple[str, int | None, int | None]:
    if isinstance(valor, datetime):
        return valor.date().isoformat(), valor.year, valor.month
    if valor in (None, ""):
        return "", None, None
    data = datetime.fromisoformat(str(valor))
    return data.date().isoformat(), data.year, data.month


def converter_total_vitima(valor: object) -> int | None:
    if valor in (None, ""):
        return None
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        return None


def montar_linha_sanitizada(
    linha: tuple[object, ...],
    posicoes: dict[str, int],
    ano_arquivo: int,
) -> dict[str, object] | None:
    uf = MAPA_UF_SUL.get(normalizar_texto(linha[posicoes["uf"]]))
    if uf not in {"RS", "SC", "PR"}:
        return None

    municipio = normalizar_texto(linha[posicoes["municipio"]])
    if municipio in {"", "NAO INFORMADO"}:
        return None

    evento_original = limpar_texto(linha[posicoes["evento"]])
    evento = EVENTOS_VIOLENCIA_GRAVE.get(normalizar_texto(evento_original))
    if evento is None:
        return None

    total_vitima = converter_total_vitima(linha[posicoes["total_vitima"]])
    if total_vitima is None:
        return None

    data_referencia, ano, mes = converter_data(linha[posicoes["data_referencia"]])
    abrangencia = normalizar_texto(linha[posicoes["abrangencia"]]) or "NAO INFORMADO"

    return {
        "ano_arquivo": ano_arquivo,
        "uf": uf,
        "municipio": municipio,
        "evento_original": evento_original,
        "evento": evento,
        "data_referencia": data_referencia,
        "ano": ano,
        "mes": mes,
        "total_vitima": total_vitima,
        "abrangencia": abrangencia,
    }


def main() -> None:
    PASTA_PROCESSADOS.mkdir(parents=True, exist_ok=True)
    PASTA_RELATORIOS.mkdir(parents=True, exist_ok=True)

    arquivos = [
        arquivo
        for arquivo in sorted(PASTA_DADOS_BRUTOS.glob("BancoVDE_*.xlsx"), key=pegar_ano_pelo_nome)
        if ANO_INICIAL <= pegar_ano_pelo_nome(arquivo) <= ANO_FINAL
    ]

    total_linhas = 0
    linhas_por_ano = Counter()
    linhas_por_uf = Counter()
    linhas_por_evento = Counter()
    municipios_por_uf: dict[str, set[str]] = {"PR": set(), "RS": set(), "SC": set()}
    base_mensal: dict[tuple[str, str, str, int, int], dict[str, object]] = {}

    with ARQUIVO_BASE_LONGA.open("w", newline="", encoding="utf-8") as arquivo_saida:
        escritor = csv.DictWriter(arquivo_saida, fieldnames=CABECALHO_BASE_LONGA)
        escritor.writeheader()

        for arquivo in arquivos:
            ano_arquivo = pegar_ano_pelo_nome(arquivo)
            print(f"Lendo {arquivo.name}...", flush=True)
            planilha = load_workbook(arquivo, read_only=True, data_only=True)
            aba = planilha[str(ano_arquivo)]
            cabecalho = [
                str(celula.value).strip() if celula.value is not None else ""
                for celula in next(aba.iter_rows(max_row=1))
            ]
            posicoes = {nome: indice for indice, nome in enumerate(cabecalho)}

            for linha in aba.iter_rows(min_row=2, values_only=True):
                dados = montar_linha_sanitizada(linha, posicoes, ano_arquivo)
                if dados is None:
                    continue

                escritor.writerow(dados)
                total_linhas += 1
                linhas_por_ano[dados["ano"]] += 1
                linhas_por_uf[dados["uf"]] += 1
                linhas_por_evento[dados["evento"]] += 1
                municipios_por_uf[str(dados["uf"])].add(str(dados["municipio"]))

                chave = (
                    str(dados["uf"]),
                    str(dados["municipio"]),
                    str(dados["data_referencia"]),
                    int(dados["ano"]),
                    int(dados["mes"]),
                )
                if chave not in base_mensal:
                    base_mensal[chave] = {
                        "uf": dados["uf"],
                        "municipio": dados["municipio"],
                        "data_referencia": dados["data_referencia"],
                        "ano": dados["ano"],
                        "mes": dados["mes"],
                        "total_violencia_grave": 0,
                        "homicidio_doloso": 0,
                        "latrocinio": 0,
                        "tentativa_homicidio": 0,
                        "lesao_corporal_seguida_morte": 0,
                        "feminicidio": 0,
                    }

                evento = str(dados["evento"])
                valor = int(dados["total_vitima"])
                base_mensal[chave][evento] = int(base_mensal[chave][evento]) + valor
                base_mensal[chave]["total_violencia_grave"] = (
                    int(base_mensal[chave]["total_violencia_grave"]) + valor
                )

            planilha.close()

    with ARQUIVO_BASE_MENSAL.open("w", newline="", encoding="utf-8") as arquivo_saida:
        escritor = csv.DictWriter(arquivo_saida, fieldnames=CABECALHO_BASE_MENSAL)
        escritor.writeheader()
        for chave in sorted(base_mensal):
            escritor.writerow(base_mensal[chave])

    with RELATORIO.open("w", encoding="utf-8") as relatorio:
        relatorio.write("Resumo da sanitizacao - violencia grave no Sul\n")
        relatorio.write(f"Periodo considerado: {ANO_INICIAL} a {ANO_FINAL}\n")
        relatorio.write(f"Arquivo base longa: {ARQUIVO_BASE_LONGA}\n")
        relatorio.write(f"Arquivo base mensal: {ARQUIVO_BASE_MENSAL}\n\n")
        relatorio.write(f"Linhas na base longa: {total_linhas}\n")
        relatorio.write(f"Linhas na base mensal agregada: {len(base_mensal)}\n\n")
        relatorio.write(f"Linhas por ano: {dict(sorted(linhas_por_ano.items()))}\n")
        relatorio.write(f"Linhas por UF: {dict(sorted(linhas_por_uf.items()))}\n")
        relatorio.write(f"Linhas por evento: {dict(sorted(linhas_por_evento.items()))}\n\n")
        for uf, municipios in sorted(municipios_por_uf.items()):
            relatorio.write(f"Municipios em {uf}: {len(municipios)}\n")
        relatorio.write(
            "\nObservacao: esta base usa contagem bruta de vitimas. "
            "Sem dados populacionais, o risco representa volume absoluto historico, "
            "nao taxa proporcional por habitante.\n"
        )

    print(f"Arquivo gerado: {ARQUIVO_BASE_LONGA}")
    print(f"Arquivo gerado: {ARQUIVO_BASE_MENSAL}")
    print(f"Relatorio gerado: {RELATORIO}")


if __name__ == "__main__":
    main()
